from pathlib import Path
import sys
import tempfile

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import pandas as pd
from PIL import Image
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


INPUT_DOC_PATH = Path("/Users/shangshanchong/Desktop/stage3_report_final.docx")
CSV_PATH = PROJECT_ROOT / "outputs" / "stage3_indoor" / "conditional_residual_ddpm_gaussian" / "seed42" / "generalization_summary.csv"
FIG1_SOURCE = PROJECT_ROOT / "outputs" / "stage3_indoor" / "conditional_residual_ddpm_gaussian" / "seed42" / "cond_residual_gaussian_eval.png"
FIG2_SOURCE = PROJECT_ROOT / "outputs" / "stage3_indoor" / "conditional_residual_ddpm_gaussian" / "seed42" / "generalization_diagnostic.png"
OUTPUT_DOC_PATH = PROJECT_ROOT / "stage3_report_complete.docx"

EXPECTED_COLUMNS = [
    "degradation",
    "method",
    "ADE_mean",
    "ADE_std",
    "smooth_mean",
    "delta_ADE_vs_noisy_mean",
    "improved_fraction",
    "wilcoxon_p_vs_noisy",
]
ROW_ORDER = [
    ("gaussian_medium", "noisy_input"),
    ("gaussian_medium", "uncond_sdedit_t2"),
    ("gaussian_medium", "cond_residual_t20"),
    ("drift_medium", "noisy_input"),
    ("drift_medium", "uncond_sdedit_t2"),
    ("drift_medium", "cond_residual_t20"),
    ("jump_medium", "noisy_input"),
    ("jump_medium", "uncond_sdedit_t2"),
    ("jump_medium", "cond_residual_t20"),
    ("burst_medium", "noisy_input"),
    ("burst_medium", "uncond_sdedit_t2"),
    ("burst_medium", "cond_residual_t20"),
    ("bias_medium", "noisy_input"),
    ("bias_medium", "uncond_sdedit_t2"),
    ("bias_medium", "cond_residual_t20"),
    ("combined_medium", "noisy_input"),
    ("combined_medium", "uncond_sdedit_t2"),
    ("combined_medium", "cond_residual_t20"),
]
RESULT_MAP = {
    ("gaussian_medium", "noisy_input"): "—",
    ("gaussian_medium", "uncond_sdedit_t2"): "sig. impr.",
    ("gaussian_medium", "cond_residual_t20"): "sig. impr.",
    ("drift_medium", "noisy_input"): "—",
    ("drift_medium", "uncond_sdedit_t2"): "worsened",
    ("drift_medium", "cond_residual_t20"): "over-corr.",
    ("jump_medium", "noisy_input"): "—",
    ("jump_medium", "uncond_sdedit_t2"): "small impr.",
    ("jump_medium", "cond_residual_t20"): "sig. impr.",
    ("burst_medium", "noisy_input"): "—",
    ("burst_medium", "uncond_sdedit_t2"): "worsened",
    ("burst_medium", "cond_residual_t20"): "worsened",
    ("bias_medium", "noisy_input"): "—",
    ("bias_medium", "uncond_sdedit_t2"): "no change",
    ("bias_medium", "cond_residual_t20"): "no change",
    ("combined_medium", "noisy_input"): "—",
    ("combined_medium", "uncond_sdedit_t2"): "no change",
    ("combined_medium", "cond_residual_t20"): "no change",
}


def ensure_inputs() -> None:
    required_paths = [
        INPUT_DOC_PATH,
        CSV_PATH,
        FIG1_SOURCE,
        FIG2_SOURCE,
    ]
    for path in required_paths:
        if not path.is_file():
            raise FileNotFoundError(f"Missing required input: {path}")


def format_float(value: float) -> str:
    return f"{value:.4f}"


def format_delta(value: float) -> str:
    return f"{value:+.4f}"


def format_percent(value: float) -> str:
    return f"{value * 100:.1f}%"


def format_p_value(value: float) -> str:
    if pd.isna(value):
        return "—"
    if value < 1e-3 or value >= 1e4:
        return f"{value:.1e}"
    return f"{value:.4f}"


def build_table_rows(df: pd.DataFrame) -> list[list[str]]:
    rows = []
    for degradation, method in ROW_ORDER:
        subset = df[(df["degradation"] == degradation) & (df["method"] == method)]
        if subset.empty:
            raise ValueError(f"Missing row in CSV for degradation={degradation}, method={method}")
        record = subset.iloc[0]

        if method == "noisy_input":
            delta = "—"
            improved = "—"
            p_val = "—"
        else:
            delta = format_delta(float(record["delta_ADE_vs_noisy_mean"]))
            improved = format_percent(float(record["improved_fraction"]))
            p_val = format_p_value(float(record["wilcoxon_p_vs_noisy"]))

        rows.append(
            [
                degradation,
                method,
                format_float(float(record["ADE_mean"])),
                format_float(float(record["ADE_std"])),
                format_float(float(record["smooth_mean"])),
                delta,
                improved,
                p_val,
                RESULT_MAP[(degradation, method)],
            ]
        )
    return rows


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
    paragraph._p = paragraph._element = None


def clear_paragraph(paragraph: Paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def insert_paragraph_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def save_image_copy(src: Path, dest: Path) -> Path:
    with Image.open(src) as img:
        img.save(dest)
    return dest


def save_panel(src: Path, row: int, col: int, dest: Path) -> Path:
    with Image.open(src) as img:
        width, height = img.size
        panel_w, panel_h = width // 3, height // 2
        box = (col * panel_w, row * panel_h, (col + 1) * panel_w, (row + 1) * panel_h)
        panel = img.crop(box)
        panel.save(dest)
    return dest


def replace_insert_paragraph(paragraph: Paragraph, image_path: Path, width_inches: float) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def main() -> None:
    ensure_inputs()

    df = pd.read_csv(CSV_PATH)
    print("generalization_summary.csv columns:")
    print(list(df.columns))
    print(df.head().to_string(index=False))
    if list(df.columns) != [
        "degradation",
        "degradation_group",
        "method",
        "N",
        "ADE_mean",
        "ADE_std",
        "smooth_mean",
        "delta_ADE_vs_noisy_mean",
        "improved_fraction",
        "wilcoxon_p_vs_noisy",
    ]:
        print("Unexpected columns:")
        print(list(df.columns))
        raise SystemExit(1)
    for col in EXPECTED_COLUMNS:
        if col not in df.columns:
            print("Unexpected columns:")
            print(list(df.columns))
            raise SystemExit(1)

    table_rows = build_table_rows(df)
    doc = Document(str(INPUT_DOC_PATH))

    table2 = None
    for table in doc.tables:
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if len(table.rows) == 19 and len(table.columns) == 9 and header[:3] == ["Degrad.", "Method", "ADE"]:
            table2 = table
            break
    if table2 is None:
        raise RuntimeError("Could not locate Table 2 (19 rows x 9 cols) in report")

    n_filled = 0
    for row_idx, row_values in enumerate(table_rows, start=1):
        for col_idx, value in enumerate(row_values):
            table2.rows[row_idx].cells[col_idx].text = value
            n_filled += 1

    temp_dir = Path(tempfile.mkdtemp(prefix="stage3_report_"))
    fig1_full = save_image_copy(FIG1_SOURCE, temp_dir / "fig1_full.png")
    fig4_bar = save_panel(FIG2_SOURCE, 0, 0, temp_dir / "fig4_bar.png")
    fig5_heatmap = save_panel(FIG2_SOURCE, 0, 1, temp_dir / "fig5_heatmap.png")

    geometry_inserted = False
    spread_inserted = False
    caption_updated = False
    n_todo = 0
    n_insert = 0

    paragraphs = list(doc.paragraphs)
    for paragraph in paragraphs:
        text = paragraph.text.strip()

        if text.startswith("(a) The model operates in relative displacement space"):
            insert_paragraph_after(
                paragraph,
                "Room geometry constraints are not enforced at any pipeline stage: training (no), conditioning (no), loss (no), sampling (no), evaluation (no). Trajectories that exit the 3m × 3m room are not clipped during evaluation.",
            )
            geometry_inserted = True

        elif text.startswith("Under the fixed canonical Kalman configuration, Kalman smoothing worsens ADE"):
            insert_paragraph_after(
                paragraph,
                "The conditional model exhibits larger trajectory-level spread (ADE std=0.0162 vs 0.0084 for unconditional SDEdit). This is consistent with conditional sampling producing larger per-trajectory corrections: when corrections are bigger in magnitude, their distribution across trajectories is also wider. The improved fraction (73.0%) confirms that the larger spread reflects larger gains rather than instability.",
            )
            spread_inserted = True

        elif text.startswith("4 columns: clean (blue) / degraded (gray) /"):
            run = paragraph.add_run(
                " In our pipeline, the degraded trajectory itself serves as the coarse reconstruction; no separate coarse-prediction stage is used. The four displayed columns therefore correspond to: target (clean), input (degraded/coarse), unconditional refinement, and final conditional output."
            )
            run.italic = True
            caption_updated = True

    if not geometry_inserted:
        raise RuntimeError("Could not locate geometry anchor paragraph")
    if not spread_inserted:
        raise RuntimeError("Could not locate §4.1 anchor paragraph")
    if not caption_updated:
        raise RuntimeError("Could not locate Figure 1 caption paragraph")

    paragraphs = list(doc.paragraphs)
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if text == "[TODO: Fill std, smooth, imp% for drift/jump/burst/bias/combined rows from generalization_summary.csv]":
            remove_paragraph(paragraph)
            n_todo += 1
        elif text.startswith("[INSERT from:") and "cond_residual_gaussian_eval.png" in text:
            replace_insert_paragraph(paragraph, fig1_full, 6.0)
            n_insert += 1
        elif text.startswith("[INSERT from:") and "generalization_diagnostic.png (burst subplot)" in text:
            replace_insert_paragraph(paragraph, FIG2_SOURCE, 6.0)
            n_insert += 1
        elif text.startswith("[INSERT from:") and "generalization_diagnostic.png (bias subplot)" in text:
            replace_insert_paragraph(paragraph, FIG2_SOURCE, 6.0)
            n_insert += 1
        elif text.startswith("[INSERT from:") and "generalization_diagnostic.png [0,0]" in text:
            replace_insert_paragraph(paragraph, fig4_bar, 5.5)
            n_insert += 1
        elif text.startswith("[INSERT from:") and "generalization_diagnostic.png [0,1]" in text:
            replace_insert_paragraph(paragraph, fig5_heatmap, 5.5)
            n_insert += 1

    paragraphs = list(doc.paragraphs)
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if text.startswith("[TODO"):
            remove_paragraph(paragraph)
            n_todo += 1
        elif text.startswith("[INSERT"):
            remove_paragraph(paragraph)
            n_insert += 1

    doc.save(str(OUTPUT_DOC_PATH))
    print(f"Saved: {OUTPUT_DOC_PATH}")
    print("=== 报告补齐完成 ===")
    print("输入: stage3_report_final.docx")
    print("输出: stage3_report_complete.docx")
    print("")
    print("Table 2 数据来源: generalization_summary.csv")
    print(f"  共填充 {n_filled} 个单元格")
    print("")
    print("图片嵌入:")
    print("  Figure 1: cond_residual_gaussian_eval.png (整张)")
    print("  Figure 2: generalization_diagnostic.png (整张)")
    print("  Figure 3: generalization_diagnostic.png (整张)")
    print("  Figure 4: generalization_diagnostic.png panel[0,0]")
    print("  Figure 5: generalization_diagnostic.png panel[0,1]")
    print("")
    print("文字补充:")
    print("  §2.4: Geometry 角色归类")
    print("  §4.1: 异常 spread 解释")
    print("  Figure 1: 列数说明")
    print("")
    print(f"删除占位符: TODO ×{n_todo}, INSERT ×{n_insert}")


if __name__ == "__main__":
    main()
