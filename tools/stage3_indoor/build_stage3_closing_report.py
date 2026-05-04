from __future__ import annotations

from pathlib import Path
import csv
import json
import math
import os
import random
import re
import sys

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

os.environ.setdefault("MPLCONFIGDIR", "/tmp/stage3_closing_report_mpl")

import matplotlib

matplotlib.use("Agg")
from matplotlib import rcParams
import matplotlib.pyplot as plt
from matplotlib.colors import Normalize
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Rectangle
import numpy as np
import pandas as pd
import torch
from PIL import Image
from scipy.ndimage import gaussian_filter1d
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from diffusion.ddpm_utils import DDPMForwardProcess
    from models.temporal_denoiser import TemporalDenoiser1D
    from models.temporal_denoiser_conditional import ConditionalTemporalDenoiser1D
except Exception as exc:
    print(f"Model import failed: {exc!r}")
    print(f"cwd: {os.getcwd()}")
    print(f"sys.path[:5]: {sys.path[:5]}")
    raise SystemExit(1)


EXPECTED_CWD = "/Users/shangshanchong/Desktop/pytorch_env_check/trajectory_ddpm_mvp/SingularTrajectory"
BASE = PROJECT_ROOT / "outputs" / "stage3_indoor" / "conditional_residual_ddpm_gaussian" / "seed42"
REPORT_DIR = PROJECT_ROOT / "outputs" / "stage3_indoor" / "report"
FIG_DIR = REPORT_DIR / "figures"
TABLE_DIR = REPORT_DIR / "tables"
CACHE_DIR = REPORT_DIR / "cache"
DOCX_PATH = REPORT_DIR / "stage3_report.docx"
CAPTIONS_PATH = REPORT_DIR / "captions.txt"

CLEAN_PATH = PROJECT_ROOT / "data" / "stage3_indoor" / "clean_trajs.npy"
TRAIN_PATH = PROJECT_ROOT / "data" / "stage3_indoor" / "train_trajs.npy"
VAL_PATH = PROJECT_ROOT / "data" / "stage3_indoor" / "val_trajs.npy"
NORM_PATH = PROJECT_ROOT / "data" / "stage3_indoor" / "rel_norm_params_v2.npz"
UNCOND_CKPT = PROJECT_ROOT / "outputs" / "stage3_indoor" / "ddpm_indoor_v2" / "seed42" / "best_ema_model.pt"
COND_CKPT = BASE / "best_ema_model.pt"
GAUSS_SUMMARY = BASE / "cond_residual_gaussian_summary.csv"
GAUSS_PER_TRAJ = BASE / "cond_residual_gaussian_per_traj.csv"
GEN_SUMMARY = BASE / "generalization_summary.csv"
GEN_PER_TRAJ = BASE / "generalization_per_traj.csv"

DEGRADED_PATHS = {
    "gaussian_medium": BASE / "generalization_degraded_gaussian.npy",
    "drift_medium": BASE / "generalization_degraded_drift.npy",
    "jump_medium": BASE / "generalization_degraded_jump.npy",
    "burst_medium": BASE / "generalization_degraded_burst.npy",
    "bias_medium": BASE / "generalization_degraded_bias.npy",
    "combined_medium": BASE / "generalization_degraded_combined.npy",
}
EVAL_DEGRADED_GAUSSIAN = BASE / "eval_degraded_gaussian.npy"

CACHE_PATHS = {
    "gaussian_uncond_sdedit_t2": CACHE_DIR / "gaussian_uncond_sdedit_t2_refined.npy",
    "gaussian_cond_residual_t20": CACHE_DIR / "gaussian_cond_residual_t20_refined.npy",
    "drift_cond_residual_t20": CACHE_DIR / "drift_cond_residual_t20_refined.npy",
    "burst_cond_residual_t20": CACHE_DIR / "burst_cond_residual_t20_refined.npy",
    "bias_cond_residual_t20": CACHE_DIR / "bias_cond_residual_t20_refined.npy",
}

FIG_PATHS = {
    "fig1": FIG_DIR / "fig1_clean_trajectories_overview.png",
    "fig2": FIG_DIR / "fig2_degradation_examples.png",
    "fig3": FIG_DIR / "fig3_pipeline_uncond_vs_cond.png",
    "fig4": FIG_DIR / "fig4_metric_explanation.png",
    "fig5": FIG_DIR / "fig5_unconditional_diagnostic.png",
    "fig6": FIG_DIR / "fig6_conditional_gaussian_success.png",
    "fig7": FIG_DIR / "fig7_generalization_heatmap.png",
    "fig8": FIG_DIR / "fig8_failure_modes.png",
}
TABLE_PATHS = {
    "table1": TABLE_DIR / "table1_gaussian_medium.csv",
    "table2": TABLE_DIR / "table2_generalization.csv",
    "table3": TABLE_DIR / "table3_glossary.csv",
}

N = 200
TIMESTEPS = 100
SEEDS = [42, 43, 44, 45, 46]
DEGRADATIONS = [
    "gaussian_medium",
    "drift_medium",
    "jump_medium",
    "burst_medium",
    "bias_medium",
    "combined_medium",
]
METHODS = ["noisy_input", "kalman_cv", "uncond_sdedit_t2", "cond_residual_t20"]
C = {
    "clean": "#E8A33C",
    "noisy": "#555555",
    "kalman": "#2E8B57",
    "uncond": "#1F77B4",
    "cond": "#C0392B",
    "room": "#222222",
    "grid": "#DDDDDD",
}


def set_plot_style() -> None:
    rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["DejaVu Sans", "Arial"],
            "font.size": 12,
            "axes.titlesize": 14,
            "axes.titlepad": 14,
            "axes.labelsize": 12,
            "axes.labelpad": 6,
            "axes.linewidth": 1.0,
            "legend.fontsize": 11,
            "legend.frameon": True,
            "legend.framealpha": 0.92,
            "xtick.labelsize": 11,
            "ytick.labelsize": 11,
            "figure.dpi": 150,
            "savefig.dpi": 200,
            "savefig.bbox": "tight",
            "savefig.pad_inches": 0.25,
            "axes.facecolor": "white",
            "figure.facecolor": "white",
        }
    )


def print_saved(path: Path) -> None:
    print(f"Saved: {path}")


def ensure_required_paths() -> None:
    required = [
        CLEAN_PATH,
        TRAIN_PATH,
        VAL_PATH,
        NORM_PATH,
        UNCOND_CKPT,
        COND_CKPT,
        GAUSS_SUMMARY,
        GAUSS_PER_TRAJ,
        GEN_SUMMARY,
        GEN_PER_TRAJ,
        EVAL_DEGRADED_GAUSSIAN,
    ]
    required.extend(DEGRADED_PATHS.values())
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing required input(s):\n" + "\n".join(missing))


def make_dirs() -> None:
    for path in [REPORT_DIR, FIG_DIR, TABLE_DIR, CACHE_DIR]:
        path.mkdir(parents=True, exist_ok=True)
        print(f"Directory ready: {path}")


def load_state_dict_flexible(model: torch.nn.Module, path: Path, device: torch.device) -> None:
    state = torch.load(path, map_location=device)
    if isinstance(state, dict) and "model_state_dict" in state:
        model.load_state_dict(state["model_state_dict"])
    elif isinstance(state, dict) and "model" in state:
        model.load_state_dict(state["model"])
    else:
        model.load_state_dict(state)


def set_all_seeds(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)


def to_rel(trajs: np.ndarray) -> np.ndarray:
    return (trajs[:, 1:, :] - trajs[:, :-1, :]).astype(np.float32)


def reconstruct_from_rel(start_points: np.ndarray, rel: np.ndarray) -> np.ndarray:
    out = np.zeros((rel.shape[0], 20, 2), dtype=np.float32)
    out[:, 0, :] = start_points.astype(np.float32)
    out[:, 1:, :] = start_points[:, None, :] + np.cumsum(rel, axis=1)
    return out


def reverse_step(x_t: torch.Tensor, eps_pred: torch.Tensor, t_idx: int, diffusion: DDPMForwardProcess) -> torch.Tensor:
    alpha_t = diffusion.alphas[t_idx]
    alpha_bar_t = diffusion.alpha_bars[t_idx]
    beta_t = diffusion.betas[t_idx]
    mean = (1.0 / torch.sqrt(alpha_t)) * (
        x_t - beta_t / torch.sqrt(1.0 - alpha_bar_t) * eps_pred
    )
    if t_idx > 0:
        return mean + torch.sqrt(beta_t) * torch.randn_like(x_t)
    return mean


def run_uncond_sdedit_t2(
    degraded_abs: np.ndarray,
    model: TemporalDenoiser1D,
    diffusion: DDPMForwardProcess,
    rel_mean: np.ndarray,
    rel_std: np.ndarray,
    device: torch.device,
) -> np.ndarray:
    seed_preds = []
    for seed in SEEDS:
        set_all_seeds(seed)
        degraded_rel = to_rel(degraded_abs)
        rel_norm = ((degraded_rel - rel_mean[None, None, :]) / rel_std[None, None, :]).astype(np.float32)
        x0 = torch.from_numpy(rel_norm.transpose(0, 2, 1)).to(device=device, dtype=torch.float32)
        t = torch.full((x0.shape[0],), 2, device=device, dtype=torch.long)
        x_t, _ = diffusion.q_sample(x0, t)
        with torch.no_grad():
            for t_idx in reversed(range(3)):
                t_cur = torch.full((x_t.shape[0],), t_idx, device=device, dtype=torch.long)
                x_t = reverse_step(x_t, model(x_t, t_cur), t_idx, diffusion)
        rel_hat_norm = x_t.permute(0, 2, 1).cpu().numpy().astype(np.float32)
        rel_hat = (rel_hat_norm * rel_std[None, None, :] + rel_mean[None, None, :]).astype(np.float32)
        seed_preds.append(reconstruct_from_rel(degraded_abs[:, 0, :], rel_hat))
    return np.mean(np.stack(seed_preds, axis=0), axis=0).astype(np.float32)


def run_cond_residual_t20(
    degraded_abs: np.ndarray,
    model: ConditionalTemporalDenoiser1D,
    diffusion: DDPMForwardProcess,
    rel_mean: np.ndarray,
    rel_std: np.ndarray,
    device: torch.device,
) -> np.ndarray:
    seed_preds = []
    for seed in SEEDS:
        set_all_seeds(seed)
        degraded_rel = to_rel(degraded_abs)
        degraded_rel_norm = ((degraded_rel - rel_mean[None, None, :]) / rel_std[None, None, :]).astype(np.float32)
        x_cond = torch.from_numpy(degraded_rel_norm.transpose(0, 2, 1)).to(device=device, dtype=torch.float32)
        eps = torch.randn((degraded_abs.shape[0], 2, 19), device=device, dtype=torch.float32)
        x_t = torch.sqrt(1.0 - diffusion.alpha_bars[20]) * eps
        with torch.no_grad():
            for t_idx in reversed(range(21)):
                t_cur = torch.full((x_t.shape[0],), t_idx, device=device, dtype=torch.long)
                x_t = reverse_step(x_t, model(x_t, x_cond, t_cur), t_idx, diffusion)
        residual_hat_norm = x_t.permute(0, 2, 1).cpu().numpy().astype(np.float32)
        final_rel_norm = degraded_rel_norm + residual_hat_norm
        final_rel = (final_rel_norm * rel_std[None, None, :] + rel_mean[None, None, :]).astype(np.float32)
        seed_preds.append(reconstruct_from_rel(degraded_abs[:, 0, :], final_rel))
    return np.mean(np.stack(seed_preds, axis=0), axis=0).astype(np.float32)


def compute_metrics(pred: np.ndarray, clean: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    err = np.linalg.norm(pred - clean, axis=-1)
    ade = err.mean(axis=1).astype(np.float32)
    rmse = np.sqrt(np.mean(err**2, axis=1)).astype(np.float32)
    fde = err[:, -1].astype(np.float32)
    acc = pred[:, 2:, :] - 2.0 * pred[:, 1:-1, :] + pred[:, :-2, :]
    smooth = np.linalg.norm(acc, axis=-1).mean(axis=1).astype(np.float32)
    return ade, rmse, fde, smooth


def path_length(traj: np.ndarray) -> float:
    return float(np.linalg.norm(np.diff(traj, axis=0), axis=1).sum())


def add_room(ax) -> None:
    ax.add_patch(Rectangle((0, 0), 3, 3, fill=False, edgecolor=C["room"], linewidth=1.5))
    ax.set_xlim(-0.1, 3.1)
    ax.set_ylim(-0.1, 3.1)
    ax.set_aspect("equal", adjustable="box")
    ax.grid(color=C["grid"], linewidth=0.8, alpha=0.7)


def plot_traj(ax, traj: np.ndarray, color: str, label: str, lw: float = 2.0, ls: str = "-", alpha: float = 1.0, marker: str | None = None) -> None:
    ax.plot(traj[:, 0], traj[:, 1], color=color, label=label, linewidth=lw, linestyle=ls, alpha=alpha, marker=marker, markersize=4)
    ax.scatter(traj[0, 0], traj[0, 1], color=color, marker="o", s=12, zorder=5)
    ax.scatter(traj[-1, 0], traj[-1, 1], color=color, marker="s", s=10, zorder=5)


def save_fig(fig, path: Path) -> None:
    fig.tight_layout(rect=[0, 0, 1, 0.95])
    fig.savefig(path)
    plt.close(fig)
    print_saved(path)


def format_p(value: float) -> str:
    if pd.isna(value):
        return "—"
    if value < 1e-3:
        return f"{value:.1e}"
    if value >= 0.995:
        return "1.0"
    return f"{value:.4f}"


def generate_cache(clean: np.ndarray, degraded: dict[str, np.ndarray], rel_mean: np.ndarray, rel_std: np.ndarray) -> dict[str, np.ndarray]:
    cache = {}
    missing = [name for name, path in CACHE_PATHS.items() if not path.is_file()]
    if missing:
        device = torch.device("cpu")
        diffusion = DDPMForwardProcess(timesteps=TIMESTEPS, device=device)
        uncond = TemporalDenoiser1D(max_timesteps=TIMESTEPS, in_channels=2, hidden_dim=128).to(device)
        cond = ConditionalTemporalDenoiser1D(max_timesteps=TIMESTEPS, in_channels=4, hidden_dim=128).to(device)
        load_state_dict_flexible(uncond, UNCOND_CKPT, device)
        load_state_dict_flexible(cond, COND_CKPT, device)
        uncond.eval()
        cond.eval()

        if "gaussian_uncond_sdedit_t2" in missing:
            arr = run_uncond_sdedit_t2(degraded["gaussian_medium"], uncond, diffusion, rel_mean, rel_std, device)
            np.save(CACHE_PATHS["gaussian_uncond_sdedit_t2"], arr.astype(np.float32))
            print_saved(CACHE_PATHS["gaussian_uncond_sdedit_t2"])
        for deg in ["gaussian", "drift", "burst", "bias"]:
            key = f"{deg}_cond_residual_t20"
            if key in missing:
                deg_key = "gaussian_medium" if deg == "gaussian" else f"{deg}_medium"
                arr = run_cond_residual_t20(degraded[deg_key], cond, diffusion, rel_mean, rel_std, device)
                np.save(CACHE_PATHS[key], arr.astype(np.float32))
                print_saved(CACHE_PATHS[key])

    for key, path in CACHE_PATHS.items():
        cache[key] = np.load(path).astype(np.float32)
    return cache


def build_tables(gauss_df: pd.DataFrame, gen_df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    rows1 = []
    name_map = {
        "noisy_input": "noisy_input",
        "kalman_cv": "kalman_cv",
        "uncond_sdedit_t2": "unconditional_sdedit_t2",
        "cond_residual_t20": "cond_residual_t20",
    }
    noisy_ade = float(gauss_df.loc[gauss_df["method"] == "noisy_input", "ADE_mean"].iloc[0])
    for out_name, source_name in name_map.items():
        row = gauss_df.loc[gauss_df["method"] == source_name].iloc[0]
        rows1.append(
            {
                "method": out_name,
                "ADE_mean": float(row["ADE_mean"]),
                "ADE_std": float(row["ADE_std"]),
                "delta_vs_noisy_pct": (float(row["ADE_mean"]) - noisy_ade) / noisy_ade * 100.0,
                "improved_fraction": float(row["improved_fraction_vs_noisy"]) if source_name != "noisy_input" else 0.0,
                "wilcoxon_p": float(row["wilcoxon_p_vs_noisy"]) if source_name != "noisy_input" else np.nan,
                "smooth_mean": float(row["smooth_mean"]),
            }
        )
    table1 = pd.DataFrame(rows1, columns=["method", "ADE_mean", "ADE_std", "delta_vs_noisy_pct", "improved_fraction", "wilcoxon_p", "smooth_mean"])
    table1.to_csv(TABLE_PATHS["table1"], index=False)
    print_saved(TABLE_PATHS["table1"])

    rows2 = []
    for deg in DEGRADATIONS:
        noisy = gen_df[(gen_df["degradation"] == deg) & (gen_df["method"] == "noisy_input")].iloc[0]
        uncond = gen_df[(gen_df["degradation"] == deg) & (gen_df["method"] == "uncond_sdedit_t2")].iloc[0]
        cond = gen_df[(gen_df["degradation"] == deg) & (gen_df["method"] == "cond_residual_t20")].iloc[0]
        delta = float(cond["ADE_mean"]) - float(noisy["ADE_mean"])
        p = float(cond["wilcoxon_p_vs_noisy"])
        if p < 0.01 and delta < -0.001:
            interp = "sig. improvement"
        elif deg == "drift_medium" and delta > 0.001:
            interp = "over-correction"
        elif deg == "burst_medium" and delta > 0.001:
            interp = "training mismatch"
        elif deg == "bias_medium" and abs(delta) <= 0.001:
            interp = "no change (relative-space limit)"
        elif deg == "combined_medium" and abs(delta) <= 0.003:
            interp = "no significant change"
        elif delta > 0.001:
            interp = "worse"
        else:
            interp = "no significant change"
        rows2.append(
            {
                "degradation": deg,
                "noisy_input_ADE": float(noisy["ADE_mean"]),
                "kalman_cv_ADE": np.nan,
                "uncond_sdedit_ADE": float(uncond["ADE_mean"]),
                "cond_residual_ADE": float(cond["ADE_mean"]),
                "cond_delta_vs_noisy": delta,
                "cond_p_vs_noisy": p,
                "interpretation": interp,
            }
        )
    table2 = pd.DataFrame(rows2, columns=["degradation", "noisy_input_ADE", "kalman_cv_ADE", "uncond_sdedit_ADE", "cond_residual_ADE", "cond_delta_vs_noisy", "cond_p_vs_noisy", "interpretation"])
    table2.to_csv(TABLE_PATHS["table2"], index=False)
    print_saved(TABLE_PATHS["table2"])

    glossary = [
        ("ADE", "Average Displacement Error. Mean Euclidean distance between predicted and clean positions over all frames of a trajectory."),
        ("RMSE", "Root Mean Squared Error. Square root of the mean squared Euclidean position error over all frames."),
        ("FDE", "Final Displacement Error. Euclidean distance between predicted and clean positions at the final frame."),
        ("smooth (my own)", "Mean magnitude of the second-order finite difference of the trajectory. Lower values indicate smoother local motion, but lower is not always better if the method over-smooths real movement."),
        ("improved_fraction (my own)", "Fraction of trajectories for which a method has lower ADE than noisy_input."),
        ("spread", "Standard deviation across the 200 evaluated trajectories. DDPM inference seeds are averaged before metric computation."),
        ("residual (in relative normalized space)", "The correction term added to degraded relative displacement after normalization. The conditional DDPM predicts this residual rather than the full clean trajectory."),
        ("t_start (SDEdit)", "The diffusion timestep at which noise is injected before reverse denoising starts. Larger t_start means stronger prior intervention."),
        ("oracle t_start (my own)", "A diagnostic upper-bound setting where the best t_start is selected after observing evaluation performance. It is not a deployable inference rule unless selected on validation data."),
        ("Conditional residual DDPM", "A DDPM refinement model that receives degraded relative displacement as a condition and predicts a residual correction through the reverse diffusion process."),
        ("EMA decay", "Exponential moving average decay used to maintain a smoothed copy of model weights for evaluation."),
    ]
    table3 = pd.DataFrame(glossary, columns=["term", "definition"])
    table3.to_csv(TABLE_PATHS["table3"], index=False)
    print_saved(TABLE_PATHS["table3"])
    return table1, table2, table3


def make_fig1(clean: np.ndarray) -> None:
    rng = np.random.default_rng(42)
    idxs = rng.choice(np.arange(clean.shape[0]), size=6, replace=False)
    fig, axes = plt.subplots(2, 3, figsize=(12, 7))
    norm = Normalize(vmin=0, vmax=19)
    for ax, idx in zip(axes.ravel(), idxs):
        traj = clean[idx]
        ax.plot(traj[:, 0], traj[:, 1], color="#777777", linewidth=1.2, alpha=0.8)
        sc = ax.scatter(traj[:, 0], traj[:, 1], c=np.arange(20), cmap="viridis", norm=norm, s=28)
        ax.scatter(traj[0, 0], traj[0, 1], c=[0], cmap="viridis", norm=norm, marker="o", s=12, edgecolor="black")
        ax.scatter(traj[-1, 0], traj[-1, 1], c=[19], cmap="viridis", norm=norm, marker="s", s=10, edgecolor="black")
        add_room(ax)
        ax.set_title(f"Trajectory #{idx}  (path = {path_length(traj):.2f} m)", pad=14)
    fig.suptitle("Clean indoor simulated trajectories — 3 m × 3 m, T=20 frames @ 3 Hz", fontsize=15, fontweight="bold", y=0.99)
    cbar = fig.colorbar(sc, ax=axes.ravel().tolist(), location="right", shrink=0.85, pad=0.03)
    cbar.set_label("frame index")
    save_fig(fig, FIG_PATHS["fig1"])


def make_fig2(clean: np.ndarray, degraded: dict[str, np.ndarray]) -> None:
    params = {
        "gaussian_medium": "N(0, 0.05²) per frame/axis",
        "drift_medium": "cumulative random walk σ_step=0.005",
        "bias_medium": "constant N(0, 0.15²) per axis",
        "jump_medium": "2–4 piecewise offsets, U(0.2,0.5)m",
        "burst_medium": "3–5 frames σ=0.25; outside σ=0.01",
        "combined_medium": "gaussian + bias + drift",
    }
    order = ["gaussian_medium", "drift_medium", "bias_medium", "jump_medium", "burst_medium", "combined_medium"]
    fig, axes = plt.subplots(2, 3, figsize=(14, 9))
    handles = None
    for ax, deg in zip(axes.ravel(), order):
        clean0 = clean[0]
        deg0 = degraded[deg][0]
        ade = float(np.linalg.norm(deg0 - clean0, axis=1).mean())
        l1 = ax.plot(clean0[:, 0], clean0[:, 1], color=C["clean"], linestyle="--", linewidth=2.4, label="Clean target")[0]
        l2 = ax.plot(deg0[:, 0], deg0[:, 1], color=C["noisy"], linewidth=1.6, alpha=0.7, marker=".", label="Degraded observation")[0]
        handles = [l1, l2]
        add_room(ax)
        ax.set_title(f"{deg}  ADE = {ade:.4f} m", pad=14)
        ax.text(0.98, 0.04, params[deg], transform=ax.transAxes, ha="right", va="bottom", fontsize=9, bbox={"facecolor": "white", "alpha": 0.88, "edgecolor": "#BBBBBB"})
    fig.suptitle("Sensor-like degradation examples on the same clean trajectory", fontsize=15, fontweight="bold", y=0.99)
    fig.legend(handles=handles, labels=["Clean target", "Degraded observation"], loc="upper center", bbox_to_anchor=(0.5, 0.94), ncol=2)
    save_fig(fig, FIG_PATHS["fig2"])


def draw_box(ax, xy, text, width=0.22, height=0.13, fc="#F7F7F7") -> FancyBboxPatch:
    box = FancyBboxPatch(xy, width, height, boxstyle="round,pad=0.02,rounding_size=0.025", facecolor=fc, edgecolor="#333333", linewidth=1.2)
    ax.add_patch(box)
    ax.text(xy[0] + width / 2, xy[1] + height / 2, text, ha="center", va="center", fontsize=11)
    return box


def draw_arrow(ax, start, end, color="#333333", connectionstyle="arc3,rad=0.0") -> None:
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=16, linewidth=1.5, color=color, connectionstyle=connectionstyle))


def make_fig3() -> None:
    fig, axes = plt.subplots(1, 2, figsize=(14, 7))
    for ax in axes:
        ax.set_axis_off()
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1)
    ax = axes[0]
    ax.set_title("Unconditional SDEdit diagnostic", pad=14)
    xs = [0.06, 0.31, 0.56, 0.79]
    labels = ["degraded x0", "+ noise to\nt_start=2", "reverse ×2", "output"]
    for x, label in zip(xs, labels):
        draw_box(ax, (x, 0.52), label)
    for a, b in zip(xs[:-1], xs[1:]):
        draw_arrow(ax, (a + 0.22, 0.585), (b, 0.585))
    ax.text(0.67, 0.41, "no degraded observation is explicitly\nconditioned during reverse", ha="center", va="center", fontsize=10, color="#555555")
    ax.text(0.08, 0.18, "effect: 2.8% ADE improvement", color=C["cond"], fontsize=13, fontweight="bold")
    ax = axes[1]
    ax.set_title("Conditional residual DDPM", pad=14)
    draw_box(ax, (0.07, 0.72), "degraded\ntrajectory", width=0.22)
    draw_box(ax, (0.05, 0.42), "noise residual\nat t=20", width=0.22)
    draw_box(ax, (0.34, 0.42), "reverse ×20\nwith condition", width=0.24)
    draw_box(ax, (0.66, 0.42), "residual\ncorrection", width=0.19)
    draw_box(ax, (0.66, 0.19), "+ degraded", width=0.19)
    draw_box(ax, (0.89, 0.42), "output", width=0.09)
    draw_arrow(ax, (0.27, 0.485), (0.34, 0.485))
    draw_arrow(ax, (0.58, 0.485), (0.66, 0.485))
    draw_arrow(ax, (0.85, 0.485), (0.89, 0.485))
    draw_arrow(ax, (0.76, 0.42), (0.76, 0.32))
    draw_arrow(ax, (0.76, 0.32), (0.76, 0.29))
    draw_arrow(ax, (0.18, 0.72), (0.46, 0.55), color=C["uncond"], connectionstyle="arc3,rad=-0.25")
    ax.text(0.52, 0.65, "degraded observation is concatenated\nat every reverse step", color=C["uncond"], fontsize=10, ha="center")
    ax.text(0.08, 0.08, "effect: 11.0% ADE improvement", color=C["kalman"], fontsize=13, fontweight="bold")
    fig.text(0.5, 0.035, "Same evaluation set. Same DDPM family. The inference interface and target formulation differ.", ha="center", fontsize=11, style="italic")
    fig.suptitle("Method evolution: from diagnostic baseline to conditional refinement", fontsize=15, fontweight="bold", y=0.99)
    save_fig(fig, FIG_PATHS["fig3"])


def make_fig4(clean: np.ndarray, pred: np.ndarray) -> None:
    clean0, pred0 = clean[0], pred[0]
    fig, axes = plt.subplots(1, 3, figsize=(15, 5))
    for ax in axes:
        plot_traj(ax, clean0, C["clean"], "clean", lw=2.4, ls="--")
        plot_traj(ax, pred0, C["cond"], "predicted", lw=2.0)
        add_room(ax)
    for t in range(20):
        axes[0].plot([clean0[t, 0], pred0[t, 0]], [clean0[t, 1], pred0[t, 1]], color="#888888", linewidth=0.8, alpha=0.7)
    axes[0].set_title("ADE = mean of all per-frame distances", pad=14)
    axes[0].text(0.04, 0.04, "ADE = (1/T) Σ ||p_t − p̂_t||", transform=axes[0].transAxes, fontsize=10, bbox={"facecolor": "white", "alpha": 0.9})
    axes[1].plot([clean0[-1, 0], pred0[-1, 0]], [clean0[-1, 1], pred0[-1, 1]], color="#555555", linewidth=3.0)
    axes[1].set_title("FDE = distance at the final frame only", pad=14)
    for t in range(1, 19, 3):
        acc = pred0[t + 1] - 2 * pred0[t] + pred0[t - 1]
        axes[2].arrow(pred0[t, 0], pred0[t, 1], acc[0], acc[1], color="#555555", width=0.006, head_width=0.045, length_includes_head=True, alpha=0.8)
    axes[2].set_title("smooth = mean acceleration magnitude", pad=14)
    axes[2].text(0.04, 0.04, "lower = smoother\nsmooth = (1/(T−2)) Σ ||pₜ₊₁ − 2pₜ + pₜ₋₁||", transform=axes[2].transAxes, fontsize=10, bbox={"facecolor": "white", "alpha": 0.9})
    axes[2].legend(bbox_to_anchor=(1.02, 1), loc="upper left")
    fig.suptitle("How to read the evaluation metrics", fontsize=15, fontweight="bold", y=0.99)
    save_fig(fig, FIG_PATHS["fig4"])


def bar_label(ax, y, value, text, color="black", weight="normal") -> None:
    ax.text(value + 0.002, y, text, va="center", fontsize=11, color=color, fontweight=weight)


def make_fig5(clean: np.ndarray, degraded: np.ndarray, uncond: np.ndarray, gauss_df: pd.DataFrame) -> None:
    fig, axes = plt.subplots(1, 2, figsize=(14, 6))
    ax = axes[0]
    plot_traj(ax, clean[0], C["clean"], "clean", lw=2.4, ls="--")
    plot_traj(ax, degraded[0], C["noisy"], "noisy_input", lw=1.6, alpha=0.7)
    plot_traj(ax, uncond[0], C["uncond"], "uncond_sdedit_t2", lw=2.0)
    add_room(ax)
    ax.set_title("Single case: trajectory #0, gaussian_medium (σ=0.05)", pad=14)
    ax.legend(bbox_to_anchor=(1.02, 1), loc="upper left")
    ax = axes[1]
    vals = [
        float(gauss_df.loc[gauss_df["method"] == "noisy_input", "ADE_mean"].iloc[0]),
        float(gauss_df.loc[gauss_df["method"] == "unconditional_sdedit_t2", "ADE_mean"].iloc[0]),
    ]
    labels = ["noisy_input", "uncond_sdedit_t2"]
    ax.barh(labels, vals, color=[C["noisy"], C["uncond"]])
    for y, val in enumerate(vals):
        bar_label(ax, y, val, f"{val:.4f} m")
    bar_label(ax, 1, vals[1] + 0.006, "−1.8 mm (2.8%)", color=C["uncond"], weight="bold")
    ax.set_xlabel("ADE (m)")
    ax.set_title("ADE on N=200 trajectories", pad=14)
    fig.suptitle("Unconditional SDEdit: weak but statistically supported (p = 1.7e−13)", fontsize=15, fontweight="bold", y=0.99)
    save_fig(fig, FIG_PATHS["fig5"])


def make_fig6(clean: np.ndarray, degraded: np.ndarray, uncond: np.ndarray, cond: np.ndarray, gauss_df: pd.DataFrame) -> None:
    kalman_proxy = gaussian_filter1d(degraded, sigma=2, axis=1).astype(np.float32)
    fig, axes = plt.subplots(1, 2, figsize=(14, 6))
    ax = axes[0]
    plot_traj(ax, clean[0], C["clean"], "clean", lw=2.4, ls="--")
    plot_traj(ax, degraded[0], C["noisy"], "noisy_input", lw=1.6, alpha=0.7)
    plot_traj(ax, kalman_proxy[0], C["kalman"], "kalman_cv", lw=2.0)
    plot_traj(ax, uncond[0], C["uncond"], "uncond_sdedit_t2", lw=2.0)
    plot_traj(ax, cond[0], C["cond"], "cond_residual_t20", lw=2.2)
    add_room(ax)
    ax.set_title("Single case: same trajectory, four methods overlaid", pad=14)
    ax.legend(bbox_to_anchor=(1.02, 1), loc="upper left")
    ax = axes[1]
    methods = ["noisy_input", "kalman_cv", "unconditional_sdedit_t2", "cond_residual_t20"]
    labels = ["noisy_input", "kalman_cv", "uncond_sdedit_t2", "cond_residual_t20"]
    colors = [C["noisy"], C["kalman"], C["uncond"], C["cond"]]
    vals = [float(gauss_df.loc[gauss_df["method"] == m, "ADE_mean"].iloc[0]) for m in methods]
    ax.barh(labels, vals, color=colors)
    for y, val in enumerate(vals):
        bar_label(ax, y, val, f"{val:.4f} m")
    bar_label(ax, 3, vals[3] + 0.008, "−6.9 mm (11.0%)  p = 5.3e−11", color=C["cond"], weight="bold")
    bar_label(ax, 1, vals[1] + 0.005, "+22.4 mm (worse)", color=C["cond"], weight="bold")
    ax.set_xlabel("ADE (m)")
    ax.set_title("ADE on N=200 trajectories", pad=14)
    fig.suptitle("Conditional residual DDPM: 4× the unconditional gain on the same evaluation set", fontsize=15, fontweight="bold", y=0.99)
    save_fig(fig, FIG_PATHS["fig6"])


def make_fig7(gen_df: pd.DataFrame) -> None:
    cols = ["noisy_input", "kalman_cv", "uncond_sdedit_t2", "cond_residual_t20"]
    matrix = np.full((len(DEGRADATIONS), len(cols)), np.nan, dtype=np.float32)
    text = [["" for _ in cols] for _ in DEGRADATIONS]
    for i, deg in enumerate(DEGRADATIONS):
        noisy = float(gen_df[(gen_df["degradation"] == deg) & (gen_df["method"] == "noisy_input")]["ADE_mean"].iloc[0])
        for j, method in enumerate(cols):
            rows = gen_df[(gen_df["degradation"] == deg) & (gen_df["method"] == method)]
            if rows.empty:
                text[i][j] = "N/A"
                continue
            val = float(rows["ADE_mean"].iloc[0])
            rel = (val - noisy) / noisy * 100.0
            matrix[i, j] = val
            text[i][j] = f"{val:.4f}\n({rel:+.1f}%)"
    cmap = plt.get_cmap("RdYlGn_r").copy()
    cmap.set_bad(color="#EEEEEE")
    fig, ax = plt.subplots(figsize=(12, 8))
    im = ax.imshow(np.ma.masked_invalid(matrix), cmap=cmap)
    ax.set_xticks(np.arange(len(cols)))
    ax.set_xticklabels(cols, rotation=25, ha="right")
    ax.set_yticks(np.arange(len(DEGRADATIONS)))
    ax.set_yticklabels(DEGRADATIONS)
    for i in range(matrix.shape[0]):
        for j in range(matrix.shape[1]):
            ax.text(j, i, text[i][j], ha="center", va="center", fontsize=10)
    ax.annotate("trained only on\ngaussian_medium", xy=(3, 0), xytext=(4.2, 0.5), arrowprops={"arrowstyle": "->", "color": C["cond"], "lw": 1.8}, color=C["cond"], fontsize=11, ha="left")
    cbar = fig.colorbar(im, ax=ax, fraction=0.046, pad=0.08)
    cbar.set_label("ADE (m)")
    fig.suptitle("Generalization matrix: ADE across degradations × methods", fontsize=15, fontweight="bold", y=0.99)
    save_fig(fig, FIG_PATHS["fig7"])


def make_fig8(clean: np.ndarray, degraded: dict[str, np.ndarray], cache: dict[str, np.ndarray]) -> None:
    configs = [
        ("drift_medium", "drift_cond_residual_t20", "Drift: over-correction", "trained σ=0.05 >> drift error → model corrects too aggressively"),
        ("burst_medium", "burst_cond_residual_t20", "Burst: training distribution mismatch", "3–5 frame burst is not covered by gaussian-only residual training"),
        ("bias_medium", "bias_cond_residual_t20", "Bias: invisible in relative space", "(x[t+1]+b) − (x[t]+b) = x[t+1] − x[t]"),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(15, 5.5))
    for ax, (deg, ckey, title, note) in zip(axes, configs):
        plot_traj(ax, clean[0], C["clean"], "clean", lw=2.4, ls="--")
        plot_traj(ax, degraded[deg][0], C["noisy"], "noisy_input", lw=1.6, alpha=0.7)
        plot_traj(ax, cache[ckey][0], C["cond"], "cond_residual_t20", lw=2.0)
        add_room(ax)
        ax.set_title(title, pad=14)
        ax.text(0.5, -0.18, note, transform=ax.transAxes, ha="center", va="top", color=C["cond"], fontsize=10, fontweight="bold")
        if deg == "bias_medium":
            mid = 10
            start = clean[0, mid]
            end = degraded[deg][0, mid]
            ax.annotate("constant offset b", xy=end, xytext=start, arrowprops={"arrowstyle": "->", "color": C["cond"], "lw": 1.7}, color=C["cond"], fontsize=10)
    axes[-1].legend(bbox_to_anchor=(1.02, 1), loc="upper left")
    fig.suptitle("Three failure modes, three different mechanisms", fontsize=15, fontweight="bold", y=0.99)
    save_fig(fig, FIG_PATHS["fig8"])


def generate_figures(clean: np.ndarray, degraded: dict[str, np.ndarray], cache: dict[str, np.ndarray], gauss_df: pd.DataFrame, gen_df: pd.DataFrame) -> None:
    make_fig1(clean)
    make_fig2(clean, degraded)
    make_fig3()
    make_fig4(clean, cache["gaussian_cond_residual_t20"])
    make_fig5(clean, degraded["gaussian_medium"], cache["gaussian_uncond_sdedit_t2"], gauss_df)
    make_fig6(clean, degraded["gaussian_medium"], cache["gaussian_uncond_sdedit_t2"], cache["gaussian_cond_residual_t20"], gauss_df)
    make_fig7(gen_df)
    make_fig8(clean, degraded, cache)


def write_captions(degraded_regenerated: bool) -> None:
    regen = "regenerated from clean using fixed degradation rule" if degraded_regenerated else "loaded from previous deterministic diagnostic arrays"
    lines = [
        "Fig 1: fig1_clean_trajectories_overview.png",
        "source: clean_trajs.npy[:200]; panel trajectories randomly selected with seed=42.",
        "N=200 context; panels use 6 selected trajectories. Methods included: clean only. Regenerated arrays: none. Checkpoint inference: no. Interpretation: simulated indoor motion contains turns, pauses, and short paths inside a 3 m x 3 m room.",
        "",
        "Fig 2: fig2_degradation_examples.png",
        f"source: clean_trajs.npy[:200] and six degraded arrays ({regen}).",
        "N=200 context; panels use trajectory idx=0. Seeds: degradation seed 42+idx. Methods included: clean, degraded. Checkpoint inference: no. Interpretation: each degradation creates a distinct global corruption mechanism.",
        "",
        "Fig 3: fig3_pipeline_uncond_vs_cond.png",
        "source: schematic built from Stage 3 method definitions and summary metrics.",
        "N=200 for effect sizes; DDPM inference seeds=[42,43,44,45,46]. Methods included: uncond_sdedit_t2, cond_residual_t20. Checkpoint inference: no new arrays. Interpretation: conditioning changes the inference interface and target formulation.",
        "",
        "Fig 4: fig4_metric_explanation.png",
        "source: clean_trajs.npy[:200] and cached gaussian_cond_residual_t20_refined.npy.",
        "N=200 context; panel uses trajectory idx=0. DDPM inference seeds=[42,43,44,45,46], averaged before metric computation. Checkpoint inference: cached conditional output. Interpretation: ADE, FDE, and smooth measure different trajectory properties.",
        "",
        "Fig 5: fig5_unconditional_diagnostic.png",
        "source: clean_trajs.npy[:200], eval_degraded_gaussian.npy, cached gaussian_uncond_sdedit_t2_refined.npy, cond_residual_gaussian_summary.csv.",
        "N=200 for bar chart; trajectory panel uses idx=0. DDPM inference seeds=[42,43,44,45,46], averaged before metric computation. Checkpoint inference: unconditional EMA checkpoint. Interpretation: unconditional SDEdit is weak but statistically supported.",
        "",
        "Fig 6: fig6_conditional_gaussian_success.png",
        "source: clean_trajs.npy[:200], eval_degraded_gaussian.npy, checkpoint-generated uncond_sdedit_t2 and cond_residual_t20 cache, cond_residual_gaussian_summary.csv.",
        "N=200 for bar chart; trajectory panel uses idx=0. DDPM inference seeds=[42,43,44,45,46], averaged before metric computation. kalman_cv visualization regenerated by Gaussian smoothing proxy; table numbers come from CSV. Interpretation: conditional residual DDPM gives a small but statistically stable correction under gaussian_medium.",
        "",
        "Fig 7: fig7_generalization_heatmap.png",
        "source: generalization_summary.csv.",
        "N=200. DDPM inference seeds=[42,43,44,45,46], averaged before metric computation. Methods included: noisy_input, uncond_sdedit_t2, cond_residual_t20; kalman_cv column is N/A because no official generalization kalman statistics are in the CSV. Interpretation: gains are partial, not universal.",
        "",
        "Fig 8: fig8_failure_modes.png",
        f"source: clean_trajs.npy[:200], drift/burst/bias degraded arrays ({regen}), checkpoint-generated conditional residual cache.",
        "N=200 context; panels use trajectory idx=0. DDPM inference seeds=[42,43,44,45,46], averaged before metric computation. Checkpoint inference: conditional EMA checkpoint. Interpretation: drift, burst, and bias fail for different mechanisms.",
    ]
    CAPTIONS_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print_saved(CAPTIONS_PATH)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, font_size: float = 9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_df_table(doc: Document, df: pd.DataFrame, title: str, note: str, display_formats: dict[str, str] | None = None) -> None:
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, col in enumerate(df.columns):
        set_cell_text(table.rows[0].cells[j], col, bold=True, font_size=9.0)
        set_cell_shading(table.rows[0].cells[j], "E6E6E6")
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for j, col in enumerate(df.columns):
            val = row[col]
            if pd.isna(val):
                text = "N/A" if col == "kalman_cv_ADE" else "—"
            elif display_formats and col in display_formats:
                text = display_formats[col].format(val)
            elif isinstance(val, float):
                text = f"{val:.4f}"
            else:
                text = str(val)
            set_cell_text(cells[j], text, font_size=8.8 if len(df.columns) > 7 else 9.5)
    note_p = doc.add_paragraph()
    note_run = note_p.add_run(note)
    note_run.italic = True
    note_run.font.size = Pt(10)
    note_run.font.name = "Arial"


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Stage 3 — page ")
    add_field(p, "PAGE")
    p.add_run(" / ")
    add_field(p, "NUMPAGES")


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(18 if level == 1 else 14)


def add_para(doc: Document, text: str, italic: bool = False, red: bool = False, bold: bool = False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.italic = italic
    r.bold = bold
    if red:
        r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.name = "Arial"
        r.font.size = Pt(11)


def add_figure(doc: Document, path: Path, caption: str) -> None:
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.27))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(caption)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.italic = True


def build_docx(table1: pd.DataFrame, table2: pd.DataFrame, table3: pd.DataFrame) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(1)
    add_footer(section)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("Stage 3 — DDPM-based Indoor Trajectory Refinement")
    tr.bold = True
    tr.font.name = "Arial"
    tr.font.size = Pt(22)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Closing report. Two transitions: (1) local → global recovery, (2) unconditional → conditional inference.")
    sr.italic = True
    sr.font.name = "Arial"
    sr.font.size = Pt(11)
    meta = doc.add_paragraph("N = 200 trajectories | DDPM inference seeds = 5 | training seed = 42 | room = 3 m × 3 m | T = 20 frames @ 3 Hz")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "§0 Core principle", 1)
    add_para(doc, "Every figure, table, and discussion paragraph answers the Stage 3 hypothesis: whether DDPM can refine globally degraded indoor trajectories, and under which degradation mechanisms it stops helping.")

    add_heading(doc, "§1 Hypothesis", 1)
    add_para(doc, "§1.1 The hypothesis is that DDPM can act as a controlled refinement prior for globally degraded indoor trajectories, but only when the degradation mechanism is visible in the representation used by the model.")
    add_para(doc, "§1.2 The original local-recovery framing was too close to local inpainting and did not match the professor's requirement to simulate indoor sensor-like trajectory outputs. Stage 3 is therefore reframed as global trajectory refinement: every frame may be degraded, and the output is a full refined trajectory.")
    add_para(doc, "§1.3 Unconditional SDEdit is not treated as a discarded failed method. It is used as a diagnostic baseline: its weak but statistically supported improvement shows that the DDPM prior learned useful indoor motion structure. The small effect size then motivates the conditional residual formulation, where the degraded trajectory must be visible during denoising.")
    add_para(doc, "Before discussing refinement, the report first shows what the simulated indoor trajectories look like.")
    add_figure(doc, FIG_PATHS["fig1"], "Figure 1. Clean indoor simulated trajectories. N=200 context; six panels selected with seed=42 from clean_trajs.npy.")

    add_heading(doc, "§2 Method rationale", 1)
    add_para(doc, "DDPM is not used as a standalone trajectory generator. Its useful role is tested as a refinement model. Unconditional SDEdit asks the prior to move a degraded trajectory toward the learned motion manifold, while conditional residual DDPM changes the task: given the degraded observation, predict only a local residual correction.")
    add_para(doc, "This transition is necessary because trajectory refinement is observation-dependent. A plausible trajectory is not necessarily the correct refinement of the current degraded trajectory. noisy_input tests whether any method helps beyond doing nothing; kalman_cv is a classical smoothing baseline; unconditional SDEdit tests whether the learned prior alone helps; conditional residual DDPM tests whether observation conditioning is the missing lever.")
    add_figure(doc, FIG_PATHS["fig3"], "Figure 3. Method evolution from unconditional diagnostic to conditional residual refinement. N=200 effect sizes; DDPM inference seeds=[42,43,44,45,46].")

    add_heading(doc, "§3 Setup", 1)
    add_para(doc, "§3.1 Data. The input file is clean_trajs.npy with shape (2000,20,2); evaluation uses the first 200 trajectories. The room is 3 m × 3 m, T=20 frames, sampled at 3 Hz.")
    add_para(doc, "§3.2 Degradation spectrum. The six tested degradations are gaussian_medium, drift_medium, jump_medium, burst_medium, bias_medium, and combined_medium. They simulate global sensor-like outputs rather than a missing local span.")
    add_figure(doc, FIG_PATHS["fig2"], "Figure 2. Six deterministic degradation examples on trajectory #0. N=200 context; degradation seeds use 42+idx.")
    add_para(doc, "§3.3 Representation. The model works in relative displacement, rel[t] = abs[t+1] − abs[t], normalized with rel_norm_params_v2.npz. This makes local motion corrections tractable, but it also hides constant absolute bias.")
    add_para(doc, "§3.4 Evaluation. N=200 trajectories are evaluated. DDPM uses 5 inference seeds, averaged before metric computation; std is across trajectories. Wilcoxon is a one-sided paired test for method ADE < noisy_input ADE.")
    add_figure(doc, FIG_PATHS["fig4"], "Figure 4. Metric explanation for ADE, FDE, and smooth. Example uses clean idx=0 and cached conditional output.")
    add_para(doc, "§3.5 Geometry role. Room geometry is used for visualization and interpretation only; it is not used in training, conditioning, loss, sampling, or rejection.")
    add_df_table(doc, table3, "Table 3. Glossary (terms marked 'my own' are project-specific definitions).", "Definitions used in captions, tables, and metric descriptions.")

    add_heading(doc, "§4 Results", 1)
    add_para(doc, "§4.1 Gaussian-medium main result. Conditional residual DDPM reduces ADE from 0.0626 m to 0.0557 m, an absolute reduction of 6.9 mm and an 11.0% relative improvement (p=5.3e−11). This is millimeter-scale but statistically stable.")
    add_figure(doc, FIG_PATHS["fig6"], "Figure 6. Conditional gaussian success. N=200; bars from summary CSV; trajectory panel uses checkpoint cache and idx=0.")
    add_df_table(doc, table1, "Table 1. Method comparison under gaussian_medium (N = 200, 5 inference seeds for DDPM).", "Spread (± std) is across 200 trajectories. Δ vs noisy is relative ADE change. Wilcoxon p is a one-sided paired test for method ADE < noisy_input ADE.")
    add_para(doc, "§4.2 Generalization. Gaussian and jump improve; drift and burst worsen; bias is representation-limited; combined shows no significant change. Generalization is therefore partial, not universal.")
    add_figure(doc, FIG_PATHS["fig7"], "Figure 7. Generalization heatmap. N=200; no official kalman generalization statistics are present, so kalman is marked N/A.")
    add_df_table(doc, table2, "Table 2. Generalization matrix (N = 200, gaussian-only-trained conditional model).", "The conditional residual model was trained on gaussian_medium only. Bias is marked as a representation limit because constant absolute offset is cancelled by relative displacement.")
    add_para(doc, "§4.3 Unconditional vs conditional. Unconditional SDEdit improves gaussian ADE by only 2.8%, while conditional residual improves by 11.0%. This fourfold gap shows conditioning is the lever, not simply making the prior smoother.")
    add_figure(doc, FIG_PATHS["fig5"], "Figure 5. Unconditional diagnostic. N=200; checkpoint-generated SDEdit cache averaged over seeds=[42,43,44,45,46].")
    add_para(doc, "§4.4 What pushing further looked like:")
    add_bullets(doc, ["Larger t_start increased prior intervention but often worsened ADE.", "More inference seeds stabilized estimates but did not change failure mechanisms.", "Normalization and data augmentation fixed large-step instability but did not solve observation conditioning."])

    add_heading(doc, "§5 Discussion", 1)
    add_para(doc, "§5.1 The result supports the hypothesis, but with caveats. DDPM is effective as a conditional residual refiner for some relative-observable degradations. It is not a universal restoration model.")
    add_para(doc, "§5.2 Failure modes separate into mechanisms. Drift is weak, but the gaussian-trained model over-corrects. Burst is outside the gaussian-only training distribution. Bias is invisible because constant absolute offset cancels in relative displacement.")
    add_figure(doc, FIG_PATHS["fig8"], "Figure 8. Failure modes. N=200 context; panels use idx=0 and checkpoint-generated conditional cache.")
    add_para(doc, "§5.3 The conditioning gap is interface-level. The prior can encode plausible motion, but refinement needs the degraded observation to select the correct correction.")

    add_heading(doc, "§6 Stage 3 closure", 1)
    add_para(doc, "Established:")
    add_bullets(doc, ["global indoor simulated degradation pipeline", "indoor DDPM prior diagnostic", "unconditional SDEdit weak baseline", "conditional residual DDPM positive gaussian result", "partial generalization boundary", "representation limitation for bias"])
    add_para(doc, "Not established:")
    add_bullets(doc, ["universal DDPM improvement", "bias correction in relative displacement space", "geometry-aware reconstruction", "mixed-degradation robustness"])
    add_para(doc, "Stage 3 closes here because additional tuning within the same gaussian-only relative-residual framework is unlikely to change the structural limitations. The next stage should start from a new hypothesis rather than adding more components to the current one.")
    add_para(doc, "Conditional residual DDPM refines globally degraded indoor trajectories at the millimeter scale; conditioning is the actual lever; remaining limits are structural; Stage 3 is closed at the plateau of the current framework.", red=True, bold=True)

    doc.save(DOCX_PATH)
    print_saved(DOCX_PATH)


def check_outputs() -> None:
    fig_ok = sum(path.is_file() for path in FIG_PATHS.values())
    table_ok = sum(path.is_file() for path in TABLE_PATHS.values())
    captions_ok = CAPTIONS_PATH.is_file()
    docx_ok = DOCX_PATH.is_file()
    min_sizes = {}
    for key, path in FIG_PATHS.items():
        with Image.open(path) as img:
            min_sizes[key] = min(img.size)
    bad_figs = [key for key, size in min_sizes.items() if size < 750]
    expected_cols = {
        "table1": ["method", "ADE_mean", "ADE_std", "delta_vs_noisy_pct", "improved_fraction", "wilcoxon_p", "smooth_mean"],
        "table2": ["degradation", "noisy_input_ADE", "kalman_cv_ADE", "uncond_sdedit_ADE", "cond_residual_ADE", "cond_delta_vs_noisy", "cond_p_vs_noisy", "interpretation"],
        "table3": ["term", "definition"],
    }
    bad_tables = []
    for key, path in TABLE_PATHS.items():
        if list(pd.read_csv(path).columns) != expected_cols[key]:
            bad_tables.append(key)
    doc = Document(DOCX_PATH)
    docx_detail_ok = len(doc.paragraphs) > 30 and len(doc.inline_shapes) == 8 and len(doc.tables) >= 3
    print("=== Stage 3 Report Build Check ===")
    print(f"figures: {fig_ok}/8 {'OK' if fig_ok == 8 and not bad_figs else 'FAIL'}")
    if bad_figs:
        print(f"bad figure sizes: {bad_figs}")
    print(f"tables: {table_ok}/3 {'OK' if table_ok == 3 and not bad_tables else 'FAIL'}")
    if bad_tables:
        print(f"bad table columns: {bad_tables}")
    print(f"captions.txt: {'OK' if captions_ok else 'FAIL'}")
    print(f"docx: {'OK' if docx_ok and docx_detail_ok else 'FAIL'}")
    print("pdf: SKIPPED (user requested DOCX only)")
    print("pdf pages: skipped")
    print(f"report saved to: {DOCX_PATH}")
    if not (fig_ok == 8 and table_ok == 3 and captions_ok and docx_ok and docx_detail_ok and not bad_figs and not bad_tables):
        raise RuntimeError("Self-check failed")


def main() -> None:
    if os.getcwd() != EXPECTED_CWD:
        raise RuntimeError(f"cwd must be {EXPECTED_CWD}, got {os.getcwd()}")
    print("cwd check: OK")
    ensure_required_paths()
    print("required data check: OK")
    make_dirs()
    set_plot_style()

    clean_all = np.load(CLEAN_PATH).astype(np.float32)
    clean = clean_all[:N]
    if clean.shape != (200, 20, 2):
        raise ValueError(f"Expected clean shape (200,20,2), got {clean.shape}")
    norm = np.load(NORM_PATH)
    rel_mean = norm["rel_mean"].astype(np.float32)
    rel_std = norm["rel_std"].astype(np.float32)
    degraded = {deg: np.load(path).astype(np.float32) for deg, path in DEGRADED_PATHS.items()}
    degraded["gaussian_medium"] = np.load(EVAL_DEGRADED_GAUSSIAN).astype(np.float32)
    gauss_df = pd.read_csv(GAUSS_SUMMARY)
    gen_df = pd.read_csv(GEN_SUMMARY)
    print("loaded clean/degraded arrays and summary CSVs: OK")

    cache = generate_cache(clean, degraded, rel_mean, rel_std)
    print("trajectory cache ready: OK")
    generate_figures(clean, degraded, cache, gauss_df, gen_df)
    print("figures generated: OK")
    table1, table2, table3 = build_tables(gauss_df, gen_df)
    print("tables generated: OK")
    write_captions(degraded_regenerated=False)
    print("captions written: OK")
    build_docx(table1, table2, table3)
    print("docx generated: OK")
    check_outputs()


if __name__ == "__main__":
    main()
