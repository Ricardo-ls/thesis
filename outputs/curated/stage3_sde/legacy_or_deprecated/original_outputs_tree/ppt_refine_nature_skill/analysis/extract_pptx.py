from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def xml_text(xml_bytes: bytes) -> str:
    root = ET.fromstring(xml_bytes)
    parts = []
    for node in root.findall(".//a:t", NS):
        if node.text:
            parts.append(node.text)
    return "\n".join(parts)


def slide_number(path: str) -> int:
    return int(re.search(r"slide(\d+)\.xml$", path).group(1))


def notes_number(path: str) -> int:
    return int(re.search(r"notesSlide(\d+)\.xml$", path).group(1))


def main() -> None:
    pptx_path = Path(__file__).resolve().parents[1] / "input" / "stage3_report_refined_dense_oral.pptx"
    prs = Presentation(str(pptx_path))

    zip_notes = {}
    with zipfile.ZipFile(pptx_path) as zf:
        for name in zf.namelist():
            if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml"):
                zip_notes[notes_number(name)] = xml_text(zf.read(name))

    rows = []
    for idx, slide in enumerate(prs.slides, start=1):
        shapes = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text.strip()
                if text:
                    shapes.append(text)
        rows.append(
            {
                "slide": idx,
                "texts": shapes,
                "notes": zip_notes.get(idx, "").strip(),
                "picture_count": sum(1 for s in slide.shapes if s.shape_type == 13),
                "shape_count": len(slide.shapes),
            }
        )

    out_json = Path(__file__).with_name("deck_extract.json")
    out_md = Path(__file__).with_name("deck_extract.md")
    out_json.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = [f"# Extracted Deck: {pptx_path.name}", ""]
    for row in rows:
        lines.append(f"## Slide {row['slide']}")
        lines.append(f"- pictures: {row['picture_count']}; shapes: {row['shape_count']}")
        lines.append("### Text")
        for text in row["texts"]:
            lines.append("```")
            lines.append(text)
            lines.append("```")
        lines.append("### Notes")
        lines.append("```")
        lines.append(row["notes"])
        lines.append("```")
        lines.append("")
    out_md.write_text("\n".join(lines), encoding="utf-8")
    print(f"slides={len(rows)}")
    print(out_md)


if __name__ == "__main__":
    main()

