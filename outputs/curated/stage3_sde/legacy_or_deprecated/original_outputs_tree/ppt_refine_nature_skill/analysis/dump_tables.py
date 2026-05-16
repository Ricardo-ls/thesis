from pathlib import Path

from pptx import Presentation

pptx_path = Path(__file__).resolve().parents[1] / "input" / "stage3_report_refined_dense_oral.pptx"
prs = Presentation(str(pptx_path))
for i, slide in enumerate(prs.slides, 1):
    for j, shape in enumerate(slide.shapes):
        if getattr(shape, "has_table", False):
            print(f"SLIDE {i} TABLE {j}")
            table = shape.table
            for row in table.rows:
                print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
            print()

