from pathlib import Path

from pptx import Presentation

pptx_path = Path(__file__).resolve().parents[1] / "input" / "stage3_report_refined_dense_oral.pptx"
prs = Presentation(str(pptx_path))
print("size", prs.slide_width, prs.slide_height)
for i, slide in enumerate(prs.slides, 1):
    print("SLIDE", i)
    for j, shape in enumerate(slide.shapes):
        text = ""
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = shape.text.replace("\n", " | ")[:100]
        print(j, shape.shape_type, int(shape.left), int(shape.top), int(shape.width), int(shape.height), text)

