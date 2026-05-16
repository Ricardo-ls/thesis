from __future__ import annotations

import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "input" / "stage3_report_refined_dense_oral.pptx"
WORK = ROOT / "output" / "stage3_report_refined_dense_oral_nature_revised_work.pptx"
OUTPUT = ROOT / "output" / "stage3_report_refined_dense_oral_nature_revised.pptx"

FIG_PRIOR = Path("SingularTrajectory/outputs/stage3_indoor/report_v2/figures/fig05_prior_sampling_check.png").resolve()
FIG_RF = Path("SingularTrajectory/outputs/stage3_indoor/receptive_field_expansion/figs/rf_expansion_comparison.png").resolve()

TITLE_UPDATES = {
    1: {
        "Keywords: DDPM · SDEdit · conditional residual model · paired trajectory recovery":
        "Keywords: DDPM · SDEdit · conditional residual model · paired full-trajectory recovery",
        "biased gap filling -> full-trajectory refinement":
        "biased gap filling -> paired full-trajectory refinement",
    },
    2: {
        "Task correction: what the old experiment really tested":
        "Task reset: why the old gap-filling result was biased",
    },
    3: {
        "2 / 13": "3 / 13",
        "Evaluation framework: paired full-trajectory recovery":
        "Evaluation framework: paired recovery, not visual plausibility",
    },
    4: {
        "3 / 13": "4 / 13",
    },
    5: {
        "5/ 13": "5 / 13",
        "Indoor DDPM prior validation before refinement":
        "Unconditional prior sanity check before refinement",
        "These checks show technical usability, not recovery ability.":
        "This verifies a usable prior; recovery is tested only in paired refinement.",
    },
    6: {
        "6/ 13": "6 / 13",
        "Controlled degradation: six error structures, one paired protocol":
        "Degradation baseline: six error structures before refinement",
        "Figs. 3–4 from the report. The same clean trajectory is degraded under all six example panels.":
        "Dataset-level baseline; values are clean vs degraded before any refinement.",
    },
    7: {
        "Unconditional SDEdit: method, reference, and sweep":
        "Unconditional SDEdit: prior-only editing with t_start",
    },
    8: {
        "Prior-only ceiling under gaussian degradation":
        "Prior-only ceiling: small gaussian gain, smoothing is not recovery",
        "Report Table 5; N = 200; spread is across trajectories after seed averaging.":
        "Report Table 5; N = 200; Linear is identity in full-trajectory setting.",
    },
    9: {
        "Conditional residual DDPM: the interface is changed, not just the model":
        "Conditional residual DDPM: condition on the observation, predict correction",
    },
    10: {
        "Matched gaussian evidence: conditioning gives the larger gain":
        "Matched gaussian result: conditioning gives the larger gain",
    },
    11: {
        "Cross-degradation test: the failures are structured":
        "Cross-degradation result: failures are structured, not random",
    },
    12: {
        "Backbone ablation: capacity is not the main bottleneck":
        "Backbone ablation: more receptive field does not fix the interface",
    },
    13: {
        "Stage 3 claim boundary and next-step implication":
        "Bounded Stage 3 claim and Stage 4 implication",
    },
}

NOTES = {
    1: [
        "Good morning. This talk reports Stage 3 of my thesis: DDPM-based refinement for indoor trajectories.",
        "The story is intentionally bounded. I first correct the original task, then test prior-only editing, and finally test whether observation conditioning changes the recovery interface.",
        "The main message is that the model needs to see the degraded observation. A stronger unconditional prior alone is not enough for this paired recovery task.",
    ],
    2: [
        "I start with the old result because it explains why the task had to change.",
        "The original setup removed frames 8 to 11, while the prefix, suffix, and both boundary points remained visible. That made the task almost a boundary-anchored bridge problem, so linear interpolation was structurally favoured.",
        "The table now includes the classical baselines as well. Linear and Savitzky-Golay are strong in the missing-span setting, while the unconditional DDPM inpainting output is not tied tightly enough to the observed path.",
        "So I read this slide as a task-interface diagnosis, not as evidence that trajectory priors are useless.",
    ],
    3: [
        "This slide fixes the evaluation language for the rest of the deck.",
        "Every row is a paired comparison: the same degraded trajectory, the same clean reference, and then one candidate output.",
        "ADE is the headline metric because the corrected task is full-trajectory recovery. RMSE, FDE, improved fraction and p-values are supporting diagnostics.",
        "The smoothness warning matters. A path can look smoother and still be farther from the clean trajectory; Kalman CV later shows exactly that.",
    ],
    4: [
        "For the corrected task, I use synthetic indoor trajectories because I need clean and degraded pairs under controlled error injection.",
        "ETH and UCY were useful for earlier prior learning, but they mostly reflect outdoor pedestrian tracks. Here the question is indoor sensor-like refinement.",
        "The pairing is the key design feature. Each clean trajectory can be degraded in known ways, and all methods are scored against the same clean reference.",
    ],
    5: [
        "This is still not a recovery result. It is only a sanity check for the unconditional indoor DDPM prior.",
        "The selected EMA checkpoint reaches its best validation loss at epoch 60. The generated samples do not show unrealistic large jumps, and the direction bias is below the internal threshold.",
        "The figure is useful as a quality-control view, but the claim stays narrow: the prior is technically usable enough to support downstream diagnostics.",
        "Actual recovery is tested only on paired clean-degraded examples, first with SDEdit and then with the conditional residual model.",
    ],
    6: [
        "This slide defines the baseline before any model tries to refine the trajectory.",
        "The six degradations probe different error structures: independent gaussian noise, smooth drift, sparse jumps, local bursts, constant bias, and a combined setting.",
        "The table is dataset-level clean-versus-degraded error for N equals 200. It is not a refinement result yet.",
        "One subtle point is drift. Its initial ADE is smaller than gaussian here, so a gaussian-trained correction can easily over-correct it later.",
    ],
    7: [
        "Now I move from prior checking to prior-only editing.",
        "SDEdit starts from the degraded trajectory, adds gaussian noise up to t_start, and denoises with the unconditional prior.",
        "Small t_start mostly preserves the input; large t_start gives the prior more freedom but also more chance to drift away from this particular observation.",
        "The best point in the sweep is t_start equals 2, with ADE moving from 0.0626 to 0.0608 metres.",
    ],
    8: [
        "This is the prior-only ceiling under matched gaussian degradation.",
        "Linear interpolation is included here for traceability, but in a full-trajectory setting with no missing mask it is effectively the identity baseline.",
        "Savitzky-Golay and Kalman show the danger of treating smoothness as recovery. Both can smooth the path, but their mean ADE worsens under this protocol.",
        "Unconditional SDEdit gives a systematic but small gain: about 2.8 percent. That motivates changing the interface rather than only polishing the prior.",
    ],
    9: [
        "This is the conceptual turn in Stage 3.",
        "SDEdit asks the prior to edit the degraded trajectory, but it does not explicitly learn what correction this observation needs.",
        "The conditional residual model receives the degraded trajectory as condition and predicts the correction residual.",
        "So the tested variable is the interface: prior-only editing versus observation-conditioned correction, under the same gaussian evaluation set.",
    ],
    10: [
        "This slide contains the strongest positive Stage 3 result.",
        "On the matched gaussian set, no refinement is 0.0626 metres ADE, unconditional SDEdit is 0.0608, and conditional residual DDPM reaches 0.0557.",
        "That is a 10.98 percent reduction versus noisy input, and an 8.41 percent relative gain over SDEdit.",
        "The wording is deliberately bounded. This supports conditioning under matched gaussian corruption; it does not prove broad sensor-error robustness.",
    ],
    11: [
        "The same gaussian-trained conditional model is then tested across other degradation types.",
        "The failures are structured. Drift worsens because the model applies too much correction to a mild smooth error. Burst worsens because the corruption is local but the model has no reliability mask.",
        "Bias is almost unchanged because a constant global offset cancels in the relative-displacement representation.",
        "So the result is not just poor generalization. It tells us what the next interface has to model: correction scale, frame reliability, and representation limits.",
    ],
    12: [
        "This slide addresses the concern that the problem may simply be model capacity.",
        "I compare the current two-block Conv1D with a four-block version whose receptive field covers the full 20-frame window.",
        "The larger backbone nearly doubles the parameter count and gives a small median relative ADE gain, but it still beats the noisy input in only one of six degradations.",
        "That points away from receptive field as the main bottleneck. The next bottleneck is observation anchoring and confidence-aware correction.",
    ],
    13: [
        "I close with a bounded claim rather than an over-claim.",
        "Stage 3 establishes five things: the old local gap task was biased, prior-only SDEdit has a small ceiling, conditioning gives the larger matched-gaussian gain, cross-degradation failures are explainable, and a wider backbone does not remove the failure regime.",
        "The main conclusion is precise: conditioning is the operative variable, but only when the corruption model, representation and test error structure are aligned.",
        "The next step is therefore sensor-anchored, confidence-aware posterior refinement, not simply a stronger unconditional prior.",
    ],
}


def replace_texts(prs: Presentation) -> None:
    for slide_idx, replacements in TITLE_UPDATES.items():
        slide = prs.slides[slide_idx - 1]
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                continue
            old = shape.text
            if old in replacements:
                shape.text = replacements[old]


def delete_shapes(slide, shape_indices: list[int]) -> None:
    for idx in sorted(shape_indices, reverse=True):
        shape = slide.shapes[idx]
        shape._element.getparent().remove(shape._element)


def style_textbox(shape, font_size=13, bold_first=False) -> None:
    for p_idx, paragraph in enumerate(shape.text_frame.paragraphs):
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(32, 43, 56)
            if p_idx == 0 and bold_first:
                run.font.bold = True


def add_figure_slides(prs: Presentation) -> None:
    # Slide 5: replace the right-hand claim boxes with the actual prior sanity-check visual.
    slide5 = prs.slides[4]
    delete_shapes(slide5, [7, 8, 9, 10])
    slide5.shapes.add_picture(str(FIG_PRIOR), Inches(8.05), Inches(1.18), width=Inches(3.7))
    box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.88), Inches(4.35), Inches(3.75), Inches(1.08))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(239, 246, 255)
    box.line.color.rgb = RGBColor(147, 197, 253)
    box.line.width = Pt(1.1)
    box.text = "Claim boundary\nUsable prior check only; paired recovery is evaluated later."
    box.text_frame.margin_left = Inches(0.18)
    box.text_frame.margin_right = Inches(0.12)
    box.text_frame.margin_top = Inches(0.10)
    style_textbox(box, font_size=12, bold_first=True)

    # Slide 12: replace the bottom prose boxes with the RF ablation figure plus a compact implication.
    slide12 = prs.slides[11]
    delete_shapes(slide12, [8, 9, 10, 11, 12, 13])
    slide12.shapes.add_picture(str(FIG_RF), Inches(0.72), Inches(3.52), width=Inches(5.55))
    box = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.65), Inches(3.78), Inches(4.70), Inches(1.42))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 251, 235)
    box.line.color.rgb = RGBColor(245, 158, 11)
    box.line.width = Pt(1.1)
    box.text = "Design implication\nThe failure regime survives full-window coverage; Stage 4 should model observation confidence and anchoring."
    box.text_frame.margin_left = Inches(0.20)
    box.text_frame.margin_right = Inches(0.16)
    box.text_frame.margin_top = Inches(0.12)
    style_textbox(box, font_size=13, bold_first=True)


def write_notes_xml(pptx_in: Path, pptx_out: Path) -> None:
    tmp = pptx_out.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(pptx_in, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("ppt/notesSlides/notesSlide") and item.filename.endswith(".xml"):
                num = int(item.filename.rsplit("notesSlide", 1)[1].split(".xml", 1)[0])
                if num in NOTES:
                    root = ET.fromstring(data)
                    for sp in root.findall(".//p:sp", NS):
                        ph = sp.find(".//p:ph", NS)
                        if ph is not None and ph.attrib.get("type") == "body":
                            tx_body = sp.find("p:txBody", NS)
                            if tx_body is not None:
                                for child in list(tx_body):
                                    if child.tag.endswith("}p"):
                                        tx_body.remove(child)
                                for paragraph_text in NOTES[num]:
                                    p = ET.SubElement(tx_body, f"{{{NS['a']}}}p")
                                    r = ET.SubElement(p, f"{{{NS['a']}}}r")
                                    ET.SubElement(r, f"{{{NS['a']}}}rPr", {"lang": "en-US", "dirty": "0"})
                                    t = ET.SubElement(r, f"{{{NS['a']}}}t")
                                    t.text = paragraph_text
                                    ET.SubElement(p, f"{{{NS['a']}}}endParaRPr", {"lang": "en-US", "dirty": "0"})
                            break
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            zout.writestr(item, data)
    shutil.move(tmp, pptx_out)


def main() -> None:
    prs = Presentation(str(INPUT))
    replace_texts(prs)
    add_figure_slides(prs)
    prs.save(WORK)
    write_notes_xml(WORK, OUTPUT)
    WORK.unlink(missing_ok=True)
    # Reopen to validate package.
    reopened = Presentation(str(OUTPUT))
    print(f"saved={OUTPUT}")
    print(f"slides={len(reopened.slides)}")


if __name__ == "__main__":
    main()

